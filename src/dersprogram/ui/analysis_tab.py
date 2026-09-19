"""Analiz: seçilen tarih aralığında öğretmen/sınıf/öğrenci/derslik bazında
ders tipine göre toplam saatleri gösterir ve Excel'e aktarır.

SADECE PROGRAMA YERLEŞMİŞ dersler sayılır - "Atanmamış Dersler"
havuzundaki bloklar hiçbir zaman hesaba katılmaz (bkz.
scheduling.get_week_view, havuzu ayrı döner).

Her tablonun EN SON sütunu "Toplam Ders"tir (kullanıcı isteği: "analiz
kısmında da toplam ders seçeneği olsun") - o kişinin/dersliğin aralıktaki
tüm ders tiplerinin toplamı. Tablo bu sütuna göre (çoktan aza) sıralı
gelir, böylece en yoğun öğretmen/öğrenci/derslik en üstte görünür.
Derslikler sekmesi, kullanıcı isteği "hangi dersliğin ne kadar
kullanıldığını takip etmek istiyoruz" için eklendi."""
from __future__ import annotations

import datetime as _dt

from PySide6.QtCore import QDate, Qt
from PySide6.QtWidgets import (
    QWidget,
    QVBoxLayout,
    QHBoxLayout,
    QLabel,
    QDateEdit,
    QPushButton,
    QTabWidget,
    QTableWidget,
    QTableWidgetItem,
    QHeaderView,
    QDialog,
    QCheckBox,
    QDialogButtonBox,
    QFileDialog,
    QMessageBox,
)

from ..db import Database, LESSON_TYPE_LABELS, TEACHERLESS_TYPES
from .. import scheduling
from . import theme

TOTAL_COLUMN_LABEL = "Toplam Ders"

# Öğretmen tablosunda gösterilmeyen ders tipleri. Deneme ve Etüt zaten
# öğretmensiz atanabilen tiplerdir (bkz. db.TEACHERLESS_TYPES); öğretmen
# analizinde bu iki sütun yer kaplamaktan başka bir şey yapmıyordu -
# kullanıcı isteği: "deneme ve etüt sekmesini de kaldıralım analiz
# partında hocalar için". Öğrenci ve derslik tablolarında duruyorlar.
TEACHER_HIDDEN_TYPES = set(TEACHERLESS_TYPES)


class AnalysisExportDialog(QDialog):
    """Excel'e aktarmadan ÖNCE hangi tabloların yazılacağını sorar
    (kullanıcı isteği: "bunu yapmadan önce sorsun öğretmen-sınıf-öğrenci
    diye seçtiklerimi eklesin excele"). Seçilen her tablo, dosyanın içinde
    AYRI BİR SEKME olur."""

    def __init__(self, options: list[tuple[str, str, int]], parent=None):
        super().__init__(parent)
        self.setWindowTitle("Excel'e Aktar")
        layout = QVBoxLayout(self)
        layout.addWidget(QLabel(
            "Hangi tablolar aktarılsın?\n"
            "Seçtiğiniz her tablo Excel dosyasında ayrı bir sekme olur."
        ))
        self.checks: dict[str, QCheckBox] = {}
        for key, label, row_count in options:
            check = QCheckBox(f"{label} ({row_count} satır)")
            check.setChecked(True)
            self.checks[key] = check
            layout.addWidget(check)

        buttons = QDialogButtonBox()
        self.export_button = buttons.addButton("Excel Oluştur", QDialogButtonBox.AcceptRole)
        self.export_button.setObjectName("primaryButton")
        buttons.addButton("Vazgeç", QDialogButtonBox.RejectRole)
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        layout.addWidget(buttons)

        for check in self.checks.values():
            check.toggled.connect(self._update_state)
        self._update_state()

    def _update_state(self) -> None:
        self.export_button.setEnabled(bool(self.selected_keys()))

    def selected_keys(self) -> list[str]:
        return [key for key, check in self.checks.items() if check.isChecked()]


def table_export_data(table: QTableWidget) -> tuple[list[str], list[list]]:
    """Bir analiz tablosunu (ekranda ne görünüyorsa aynen, aynı sırayla)
    başlık satırı + veri satırları olarak döner. Sayısal hücreler Excel'de
    de sayı olarak yazılsın diye int'e çevrilir."""
    headers = [table.horizontalHeaderItem(c).text() for c in range(table.columnCount())]
    rows: list[list] = []
    for r in range(table.rowCount()):
        row: list = []
        for c in range(table.columnCount()):
            item = table.item(r, c)
            text = item.text() if item is not None else ""
            if c == 0:
                row.append(text)
            else:
                row.append(int(text) if text.lstrip("-").isdigit() else text)
        rows.append(row)
    return headers, rows


def write_analysis_workbook(
    path: str, sheets: list[tuple[str, list[str], list[list]]],
    start: _dt.date, end: _dt.date,
) -> None:
    """Seçilen analiz tablolarını tek bir .xlsx dosyasına, her tabloyu
    AYRI SEKME olarak yazar. openpyxl zaten bir bağımlılık (öğrenci
    içe aktarma da kullanıyor), ek bir kurulum gerekmez."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)
    title_font = Font(bold=True, size=13)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E5F")
    total_font = Font(bold=True)
    center = Alignment(horizontal="center")

    period = f"{start.strftime('%d.%m.%Y')} - {end.strftime('%d.%m.%Y')}"
    for sheet_name, headers, rows in sheets:
        sheet = workbook.create_sheet(sheet_name[:31])
        sheet["A1"] = f"{sheet_name} - Ders Analizi"
        sheet["A1"].font = title_font
        sheet["A2"] = f"Tarih aralığı: {period} (sadece programa yerleşmiş dersler)"
        sheet.append([])
        sheet.append(headers)
        header_row = sheet.max_row
        for cell in sheet[header_row]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
        for row in rows:
            sheet.append(row)
            for cell in sheet[sheet.max_row][1:]:
                cell.alignment = center
            sheet.cell(row=sheet.max_row, column=len(headers)).font = total_font

        # Sütun genişlikleri içeriğe göre; başlık satırı sabitlenir ki
        # uzun listelerde kaydırırken başlıklar ekranda kalsın.
        for col, header in enumerate(headers, start=1):
            longest = max([len(str(header))] + [len(str(row[col - 1])) for row in rows] or [0])
            sheet.column_dimensions[get_column_letter(col)].width = min(42, max(10, longest + 4))
        sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
        sheet.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(len(headers))}{header_row + len(rows)}"
        )
    workbook.save(path)


class AnalysisTab(QWidget):
    def __init__(self, db: Database):
        super().__init__()
        self.db = db

        layout = QVBoxLayout(self)

        date_row = QHBoxLayout()
        today = QDate.currentDate()
        date_row.addWidget(QLabel("Başlangıç (dahil):"))
        self.start_edit = QDateEdit(today.addDays(-27))
        self.start_edit.setCalendarPopup(True)
        self.start_edit.setToolTip("Bu tarih hesaba DAHİL edilir.")
        date_row.addWidget(self.start_edit)
        date_row.addWidget(QLabel("Bitiş (dahil):"))
        self.end_edit = QDateEdit(today)
        self.end_edit.setCalendarPopup(True)
        self.end_edit.setToolTip("Bu tarih hesaba DAHİL edilir.")
        date_row.addWidget(self.end_edit)
        self.refresh_button = QPushButton("Hesapla")
        self.refresh_button.setObjectName("primaryButton")
        date_row.addWidget(self.refresh_button)
        self.export_button = QPushButton("  Excel'e Aktar")
        self.export_button.setObjectName("outlineButton")
        self.export_button.setIcon(theme.icon(theme.NAV_ICONS["document"], theme.ACCENT_HOVER))
        self.export_button.setToolTip(
            "Seçtiğiniz tabloları tek bir Excel dosyasına, her biri ayrı sekme olacak şekilde aktarır."
        )
        self.export_button.clicked.connect(self.handle_export_excel)
        date_row.addWidget(self.export_button)
        date_row.addStretch()
        layout.addLayout(date_row)

        self.tabs = QTabWidget()
        self.teacher_table = self._make_table(hidden_types=TEACHER_HIDDEN_TYPES)
        self.class_table = self._make_table()
        self.student_table = self._make_table()
        self.room_table = self._make_table()
        # (anahtar, başlık, tablo, hangi alana göre gruplanacağı, kaynak)
        # - hem ekran sekmeleri hem Excel sekmeleri bu tek listeden gelir.
        self._sections = [
            ("teacher", "Öğretmenler", self.teacher_table, "teacher_id", self.db.list_teachers),
            ("class", "Sınıflar", self.class_table, "class_group_id", self.db.list_class_groups),
            ("student", "Öğrenciler", self.student_table, "student_id", self.db.list_students),
            ("room", "Derslikler", self.room_table, "room_id", lambda: self.db.list_rows("rooms")),
        ]
        for _key, label, table, _field, _source in self._sections:
            self.tabs.addTab(table, label)
        layout.addWidget(self.tabs, 1)

        self.refresh_button.clicked.connect(self.refresh)
        self.refresh()

    def _make_table(self, hidden_types: set[str] | None = None) -> QTableWidget:
        table = QTableWidget()
        table.setEditTriggers(QTableWidget.NoEditTriggers)
        types = [t for t in LESSON_TYPE_LABELS if t not in (hidden_types or set())]
        # Hangi tiplerin gösterildiği tabloyla birlikte taşınır - _fill_table
        # hem sütunları hem TOPLAM'ı bu listeye göre doldurur, böylece
        # "Toplam Ders" her zaman görünen sütunların toplamına eşit kalır.
        table.lesson_types = types
        columns = ["Ad"] + [LESSON_TYPE_LABELS[t] for t in types] + [TOTAL_COLUMN_LABEL]
        table.setColumnCount(len(columns))
        table.setHorizontalHeaderLabels(columns)
        table.horizontalHeader().setSectionResizeMode(0, QHeaderView.Stretch)
        # Sütun başlıkları ("Öğrenci Koçluk", "Soru Çözümü") varsayılan
        # genişlikte kırpılıyordu. ResizeToContents KULLANILMAZ - o mod
        # satır sayısı büyüdükçe (400 öğrenci) tüm hücreleri yeniden
        # ölçtüğü için ekranı yavaşlatıyor; başlık metni bir kez ölçülüp
        # sabit genişlik veriliyor.
        metrics = table.fontMetrics()
        for col, text in enumerate(columns):
            if col == 0:
                continue
            table.setColumnWidth(col, metrics.horizontalAdvance(text) + 26)
        table.setToolTip(
            "Son sütun (Toplam Ders), seçilen tarih aralığındaki tüm ders "
            "tiplerinin toplamıdır. Tablo bu sütuna göre sıralıdır."
        )
        return table

    def refresh(self) -> None:
        start = self.start_edit.date().toPython()
        end = self.end_edit.date().toPython()
        if start > end:
            start, end = end, start

        for _key, _label, table, field, source in self._sections:
            self._fill_table(table, source(), start, end, field)

    def _fill_table(self, table: QTableWidget, entities, start: _dt.date, end: _dt.date, key: str) -> None:
        types = table.lesson_types
        # Tüm satırların toplamı TEK geçişte hesaplanır (eskiden satır
        # başına bir kez aralığın tamamı yeniden kuruluyordu), sonra TOPLAM
        # DERS'e göre (çoktan aza) sıralanır - en yoğun olan en üstte.
        by_id = scheduling.summarize_hours_range_by(self.db, start, end, key)
        empty = {t: 0 for t in types}
        rows = []
        for entity in entities:
            totals = by_id.get(entity["id"], empty)
            rows.append((entity["name"], totals, sum(totals.get(t, 0) for t in types)))
        rows.sort(key=lambda row: (-row[2], row[0]))

        table.setRowCount(len(rows))
        total_col = len(types) + 1
        for r, (name, totals, grand_total) in enumerate(rows):
            table.setItem(r, 0, QTableWidgetItem(name))
            for c, t in enumerate(types, start=1):
                item = QTableWidgetItem(str(totals.get(t, 0)))
                item.setTextAlignment(Qt.AlignCenter)
                table.setItem(r, c, item)
            total_item = QTableWidgetItem(str(grand_total))
            total_item.setTextAlignment(Qt.AlignCenter)
            font = total_item.font()
            font.setBold(True)
            total_item.setFont(font)
            table.setItem(r, total_col, total_item)

    # ---------- Excel ----------
    def handle_export_excel(self) -> None:
        """Önce hangi tabloların aktarılacağını sorar, sonra seçilenleri
        tek bir Excel dosyasına her biri AYRI SEKME olacak şekilde yazar."""
        # Ekrandaki sayılar tarih aralığıyla uyumsuz kalmasın diye önce
        # yeniden hesaplanır - kullanıcı tarihi değiştirip "Hesapla"ya
        # basmadan doğrudan aktarmaya kalkabilir.
        self.refresh()

        options = [(key, label, table.rowCount()) for key, label, table, _f, _s in self._sections]
        dialog = AnalysisExportDialog(options, self)
        if dialog.exec() != QDialog.Accepted:
            return
        chosen = set(dialog.selected_keys())
        if not chosen:
            return

        start = self.start_edit.date().toPython()
        end = self.end_edit.date().toPython()
        if start > end:
            start, end = end, start

        default_name = f"analiz_{start:%Y%m%d}_{end:%Y%m%d}.xlsx"
        path, _filter = QFileDialog.getSaveFileName(
            self, "Excel Olarak Kaydet", default_name, "Excel Dosyası (*.xlsx)",
        )
        if not path:
            return
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"

        sheets = []
        for key, label, table, _field, _source in self._sections:
            if key not in chosen:
                continue
            headers, rows = table_export_data(table)
            sheets.append((label, headers, rows))

        try:
            write_analysis_workbook(path, sheets, start, end)
        except OSError as error:
            # En sık sebep: dosya Excel'de açık olduğu için üzerine yazılamıyor.
            QMessageBox.warning(
                self, "Kaydedilemedi",
                f"Excel dosyası kaydedilemedi:\n{error}\n\n"
                "Dosya başka bir programda açıksa kapatıp tekrar deneyin.",
            )
            return
        QMessageBox.information(
            self, "Tamamlandı",
            f"{len(sheets)} sekme şu dosyaya kaydedildi:\n{path}",
        )
